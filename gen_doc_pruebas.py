import docx
from docx.shared import Pt, RGBColor

doc = docx.Document()

# Styles
doc.add_heading('Plan de Pruebas: Onboarding de Clientes (Core Bancario)', 0)
doc.add_paragraph('Este documento contiene los escenarios de prueba exhaustivos (Happy Path y Negative/Boundary Testing) para los modulos de Clientes, Cuentas, Usuarios y Autenticacion del sistema de Onboarding. Guíate con estos payloads y resultados esperados para capturar tus evidencias en Postman.')

def add_test_case(code, title, expected_status, request_method_url, request_body, expected_response, note=None):
    p = doc.add_paragraph()
    run = p.add_run(f'{code} - {title}')
    run.bold = True
    run.font.size = Pt(12)
    
    p2 = doc.add_paragraph()
    p2.add_run('Resultado esperado: ').bold = True
    p2.add_run(str(expected_status))
    
    p3 = doc.add_paragraph()
    p3.add_run('Petición:').bold = True
    doc.add_paragraph(request_method_url)
    if request_body:
        doc.add_paragraph(request_body)
        
    p4 = doc.add_paragraph()
    p4.add_run('Respuesta esperada:').bold = True
    doc.add_paragraph(expected_response)
    
    if note:
        p5 = doc.add_paragraph()
        r = p5.add_run(f'Nota: {note}')
        r.italic = True
    
    doc.add_paragraph('-' * 40)

# 1. Autenticación
doc.add_heading('1. Autenticacion (Auth)', level=1)
add_test_case('AUTH-01', 'Login Exitoso', '200 OK', 'POST /auth/login', 
'''{
  "correo": "carlos.gomez@gmail.com",
  "password": "Password123!"
}''',
'''{
  "token": "eyJhbGciOiJIUzI1NiJ9..."
}''')

add_test_case('AUTH-02', 'Login Fallido por Credenciales Incorrectas', '403 Forbidden / 401 Unauthorized', 'POST /auth/login', 
'''{
  "correo": "carlos.gomez@gmail.com",
  "password": "ClaveIncorrecta!"
}''',
'''{
  "status": 403,
  "error": "Forbidden"
}''', 'La capa de Spring Security debe rechazar la solicitud por credenciales incorrectas.')

# 2. Clientes
doc.add_heading('2. Clientes (Onboarding)', level=1)
add_test_case('CLI-01', 'Registro de Cliente Exitoso (Happy Path)', '201 Created', 'POST /clientes',
'''{
  "nombre": "Carlos",
  "segundoNombre": "Alberto",
  "apellidoPaterno": "Gomez",
  "apellidoMaterno": "Perez",
  "fechaNacimiento": "1990-05-15",
  "curp": "GOPC900515HMCNRA01",
  "rfc": "GOPC900515XYZ",
  "sexo": "MASCULINO",
  "nacionalidad": "MEXICANA",
  "estadoCivil": "SOLTERO",
  "correoElectronico": "carlos.gomez@gmail.com",
  "ladaMovil": "52",
  "telefonoMovil": "5512345678",
  "calle": "Av. Insurgentes",
  "numeroExterior": "123",
  "colonia": "Roma Norte",
  "municipio": "Cuauhtemoc",
  "estado": "CIUDAD_DE_MEXICO",
  "codigoPostal": "06700",
  "pais": "MEXICO",
  "ocupacion": "Ingeniero",
  "empresa": "Tech Solutions",
  "ingresoMensual": 35000.00,
  "password": "Password123!"
}''',
'''{
  "id": 1,
  "nombre": "Carlos",
  "apellidoPaterno": "Gomez",
  ... 
}''', 'Este registro dispara la creación automática de la Cuenta Bancaria y del Usuario asociado.')

add_test_case('CLI-02', 'Fallo por Cliente Menor de Edad (Boundary Testing)', '400 / 500', 'POST /clientes',
'''{
  "fechaNacimiento": "2015-05-15",
  ... (resto de datos correctos)
}''',
'''{
  "mensaje": "El cliente debe ser mayor de 18 años"
}''')

add_test_case('CLI-03', 'Fallo por CURP o Correo Duplicado (Race / Unique Constraint)', '400 / 500', 'POST /clientes',
'''(Mismo payload del CLI-01)''',
'''{
  "mensaje": "Ya existe un cliente con la CURP proporcionada"
}''')

add_test_case('CLI-04', 'Consulta de Cliente por ID', '200 OK', 'GET /clientes/1\nAuthorization: Bearer <JWT>', '',
'''{
  "id": 1,
  "nombre": "Carlos",
  "curp": "GOPC900515HMCNRA01"...
}''')

add_test_case('CLI-05', 'Actualización Parcial del Cliente', '200 OK', 'PUT /clientes/1\nAuthorization: Bearer <JWT>',
'''{
  "ingresoMensual": 45000.00,
  "estadoCivil": "CASADO"
}''',
'''{
  "id": 1,
  "ingresoMensual": 45000.00,
  "estadoCivil": "CASADO"...
}''', 'La CURP y RFC no se alteran aunque se envíen porque no están permitidos.')

add_test_case('CLI-06', 'Baja Lógica del Cliente', '204 No Content / 200 OK', 'DELETE /clientes/1\nAuthorization: Bearer <JWT>', '',
'''(Sin cuerpo)''', 'Inactiva al cliente y cambia automáticamente a su Usuario a Inactivo.')

# 3. Cuentas
doc.add_heading('3. Cuentas Bancarias', level=1)
add_test_case('CTA-01', 'Consulta de Cuenta por Número', '200 OK', 'GET /cuentas/0000001234\nAuthorization: Bearer <JWT>', '',
'''{
  "id": 1,
  "numeroCuenta": "0000001234",
  "saldo": 1000.00,
  "estatus": "ACTIVA",
  "cliente": { "id": 1 }
}''')

add_test_case('CTA-02', 'Búsqueda de Cuentas por ID Cliente', '200 OK', 'GET /cuentas?clienteId=1\nAuthorization: Bearer <JWT>', '',
'''[
  {
    "id": 1,
    "numeroCuenta": "0000001234",
    "estatus": "ACTIVA"
  }
]''')

add_test_case('CTA-03', 'Actualizar Estatus de la Cuenta', '200 OK', 'PUT /cuentas/0000001234?estatus=BLOQUEADA\nAuthorization: Bearer <JWT>', '',
'''{
  "id": 1,
  "numeroCuenta": "0000001234",
  "estatus": "BLOQUEADA"
}''')

# 4. Transversales
doc.add_heading('4. Escenarios Transversales de Seguridad y Formato', level=1)
add_test_case('TRA-01', 'Acceso sin Token a Ruta Protegida', '403 / 401', 'GET /clientes', '',
'''{
  "status": 403,
  "error": "Forbidden"
}''')

add_test_case('TRA-02', 'Payload malformado (Tipos de datos inválidos)', '400 Bad Request', 'POST /clientes',
'''{
  "ingresoMensual": "NO_SOY_UN_NUMERO"
}''',
'''{
  "status": 400,
  "error": "Bad Request"
}''')

path = r'C:\Users\bhleo\Downloads\Plan_Pruebas_Onboarding_Completo.docx'
doc.save(path)
print('Document created successfully at:', path)
